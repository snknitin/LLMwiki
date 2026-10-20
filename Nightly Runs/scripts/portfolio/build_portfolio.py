"""Build Nightly Runs/Project Portfolio.md and Project Portfolio.xlsx from portfolio.csv plus the live project repos.

portfolio.csv (next to this script) holds the judgment fields: status, best use, DCC group, fit scores, rationale, notes.
Purpose and stack are read from each project's BUILD_SPEC.md (§1, §3); open-question counts from its §6 (same rules as
handoff_questions.py, so they match REPORT.md and the Agents board). Notes typed into the spreadsheet's "Your notes"
column are copied back into portfolio.csv before rebuilding, so edits survive regeneration.

Usage:  uv run --no-project --with openpyxl python build_portfolio.py [--projects E:/GIT_ROOT/Projects]
"""
import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from handoff_questions import groups_for

HERE = Path(__file__).resolve().parent
NIGHTLY = HERE.parents[1]
CSV_PATH = HERE / "portfolio.csv"
XLSX_PATH = NIGHTLY / "Project Portfolio.xlsx"
MD_PATH = NIGHTLY / "Project Portfolio.md"
INACTIVE = ("Merged → feature", "Parked (archived)", "Retired (archived)", "Third-party")

STACK_KW = [("SvelteKit", r"SvelteKit"), ("Next.js", r"Next\.js"), ("Expo", r"\bExpo\b"), ("Kotlin/Compose", r"Kotlin|Jetpack Compose"),
            ("Tauri", r"Tauri"), ("Astro", r"\bAstro\b"), ("Streamlit", r"Streamlit"), ("Gradio", r"Gradio"), ("FastAPI", r"FastAPI"),
            ("HTMX", r"HTMX"), ("Jinja", r"Jinja"), ("n8n", r"\bn8n\b"), ("OR-Tools", r"OR-Tools|CP-SAT"),
            ("sentence-transformers", r"sentence-transformers"), ("FAISS", r"FAISS"), ("Polars", r"Polars"), ("DuckDB", r"DuckDB"),
            ("Playwright", r"Playwright"), ("LangGraph", r"LangGraph"), ("DSPy", r"DSPy"), ("XState", r"XState"),
            ("Three.js/WebGL", r"Three\.js|WebGL"), ("ComfyUI", r"ComfyUI"), ("Supabase", r"Supabase"), ("Cloudflare", r"Cloudflare"),
            ("PyTorch", r"PyTorch|\btorch\b"), ("FFmpeg", r"FFmpeg"), ("SQLite", r"SQLite"), ("Telegram", r"Telegram"),
            ("Chrome MV3", r"Manifest V3|\bMV3\b"), ("React", r"\bReact\b"), ("Qt", r"\bQt\b|PySide"), ("ARCore", r"ARCore"),
            ("Gymnasium", r"Gymnasium"), ("k6", r"\bk6\b"), ("Docker", r"Docker")]


def spec_facts(spec: str):
    m1 = re.search(r"^## 1\..*?$(.*?)(?=^## 2\.)", spec, re.S | re.M)
    purpose = re.split(r"\s*Not in scope", re.sub(r"\s+", " ", m1.group(1)).strip())[0] if m1 else ""
    purpose = re.sub(r"^>\s*\*\*Status.*?\*\*\s*", "", purpose)
    m3 = re.search(r"^## 3\..*?(?=^## 4\.)", spec, re.S | re.M)
    stack = ", ".join(n for n, rx in STACK_KW if m3 and re.search(rx, m3.group(0)))
    g = groups_for(spec)
    return purpose, stack, (len(g["blocking"]), len(g["question"]), len(g["proposal"]))


def sync_notes_from_xlsx(rows):
    if not XLSX_PATH.exists():
        return 0
    ws = load_workbook(XLSX_PATH)["Portfolio"]
    head = [c.value for c in ws[1]]
    if "Your notes" not in head:
        return 0
    ni, changed = head.index("Your notes"), 0
    xl = {r[0].value: (r[ni].value or "") for r in ws.iter_rows(min_row=2) if r[0].value}
    for r in rows:
        if xl.get(r["project"], "") and xl[r["project"]] != r["notes"]:
            r["notes"], changed = xl[r["project"]], changed + 1
    return changed


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--projects", default="E:/GIT_ROOT/Projects")
    root = Path(ap.parse_args().projects)

    with open(CSV_PATH, encoding="utf-8", newline="") as f:
        rows = list(csv.DictReader(f))
    fields = list(rows[0].keys())
    synced = sync_notes_from_xlsx(rows)
    if synced:
        with open(CSV_PATH, "w", encoding="utf-8", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields)
            w.writeheader()
            w.writerows(rows)

    on_disk = {p.name for p in root.iterdir() if (p / "README.md").exists()}
    missing = sorted(on_disk - {r["project"] for r in rows})
    for r in rows:
        spec = root / r["project"] / "BUILD_SPEC.md"
        purpose, stack, qs = spec_facts(spec.read_text(encoding="utf-8")) if spec.exists() else ("", "", None)
        r["purpose"] = r["purpose_override"] or purpose
        r["stack"] = "—" if r["status"] in INACTIVE and r["status"] != "Merged → feature" else (r["stack_override"] or stack or "—")
        r["open"] = "—" if (qs is None or r["status"] in INACTIVE or r["status"] in ("Brownfield (live)", "Built elsewhere", "Kept (ideas ported)")) \
            else f"{qs[0]} blocking · {qs[1]} questions · {qs[2]} proposals"
        r["repo"] = "https://github.com/Panniantong/agent-reach" if r["project"] == "agent-reach" else f"https://github.com/snknitin/{r['project']}"
    rows.sort(key=lambda r: (r["category"], r["project"]))

    write_xlsx(rows)
    write_md(rows, missing)
    print(f"rows={len(rows)} notes_synced={synced} projects_missing_from_csv={missing}")


def write_xlsx(rows):
    FONT = "Arial"
    hdr_fill, hdr_font, body = PatternFill("solid", fgColor="1F3A5F"), Font(name=FONT, bold=True, color="FFFFFF", size=10), Font(name=FONT, size=10)
    border = Border(bottom=Side(style="thin", color="D0D7E2"))
    tint = {"Brownfield (live)": "E3F2E1", "Greenfield · host": "E1ECF7", "Built elsewhere": "FFF4D6", "Kept (ideas ported)": "FFF4D6"}
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # Excel recomputes the count formulas on open
    ws = wb.active
    ws.title = "Portfolio"
    cols = [("Project", "project", 30), ("Status", "status", 19), ("Merged / parked into", "merged_or_parked_into", 30),
            ("One-line purpose", "purpose", 60), ("Category", "category", 15), ("Form factor", "form_factor", 26),
            ("Primary stack (from spec §3)", "stack", 34), ("Runs on", "runs_on", 24), ("Best use", "best_use", 18),
            ("DCC group", "dcc_group", 15), ("Best path", "best_path", 22), ("Open-source fit", "open_source_fit", 13),
            ("Commercial fit", "commercial_fit", 13), ("Rationale", "rationale", 55), ("GitHub repo", "repo", 44),
            ("Your notes", "notes", 30), ("Open questions (REPORT.md)", "open", 30)]
    ws.append([c[0] for c in cols])
    for i, (_, _, w) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font, c.alignment = hdr_fill, hdr_font, Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([r[k] for _, k, _ in cols])
        rr = ws.max_row
        fill = PatternFill("solid", fgColor=tint.get(r["status"], "F2F2F2" if r["status"] in INACTIVE else "FFFFFF"))
        for c in ws[rr]:
            c.font, c.border, c.fill, c.alignment = body, border, fill, Alignment(wrap_text=True, vertical="top")
        ws.cell(rr, 15).hyperlink = r["repo"]
        ws.cell(rr, 15).font = Font(name=FONT, size=10, color="0563C1", underline="single")
        ws.cell(rr, 16).fill = PatternFill("solid", fgColor="FFFF00")
    last = ws.max_row
    ws.freeze_panes, ws.auto_filter.ref = "B2", f"A1:Q{last}"
    for col in ("L", "M"):
        dv = DataValidation(type="list", formula1='"High,Medium,Low,—"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{last}")

    mg = wb.create_sheet("Merges")
    mg.append(["Project", "Status", "Now lives in", "Note"])
    for r in rows:
        if r["status"] in INACTIVE or r["status"] in ("Built elsewhere", "Kept (ideas ported)"):
            mg.append([r["project"], r["status"], r["merged_or_parked_into"] or "—", r["rationale"]])
    for i, w in enumerate([34, 20, 36, 70], 1):
        mg.column_dimensions[get_column_letter(i)].width = w

    dg = wb.create_sheet("DCC Groups")
    dg.append(["DCC group", "Dashboards in group (live count)", "Members"])
    for g in ["Money", "Career", "Learning", "Signals", "Publishing", "Ventures", "Client services", "Personal", "Dev", "Studio", "Labs"]:
        members = sorted(r["project"] for r in rows if r["dcc_group"] == g and r["best_use"].startswith(("DCC", "Standalone / DCC")))
        n = dg.max_row + 1
        dg.append([g, f'=COUNTIFS(Portfolio!$J$2:$J${last},A{n},Portfolio!$I$2:$I${last},"DCC*")+COUNTIFS(Portfolio!$J$2:$J${last},A{n},Portfolio!$I$2:$I${last},"Standalone / DCC tab")', ", ".join(members)])
    for i, w in enumerate([18, 30, 110], 1):
        dg.column_dimensions[get_column_letter(i)].width = w

    sm = wb.create_sheet("Summary")
    sm["A1"], sm["B1"] = "Projects in portfolio", f"=COUNTA(Portfolio!A2:A{last})"
    at = 3
    for title, col, values in (("By status", "B", sorted({r["status"] for r in rows})), ("By best use", "I", sorted({r["best_use"] for r in rows})),
                               ("By commercial fit", "M", ["High", "Medium", "Low"]), ("By open-source fit", "L", ["High", "Medium", "Low"])):
        sm.cell(at, 1, title).font = Font(name=FONT, bold=True, size=11)
        for i, v in enumerate(values, 1):
            sm.cell(at + i, 1, v)
            sm.cell(at + i, 2, f"=COUNTIF(Portfolio!${col}$2:${col}${last},A{at + i})")
        sm.cell(at + len(values) + 1, 1, "Total").font = Font(name=FONT, bold=True, size=10)
        sm.cell(at + len(values) + 1, 2, f"=SUM(B{at + 1}:B{at + len(values)})")
        at += len(values) + 3
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width = 38, 12

    lg = wb.create_sheet("Legend")
    lg.append(["Item", "Meaning"])
    for a, b in (("Source of truth", "scripts/portfolio/portfolio.csv. This workbook and Project Portfolio.md are generated by build_portfolio.py."),
                 ("Your notes (yellow)", "Type in column P; the next build copies them into portfolio.csv. Edits to other columns are overwritten — change those in portfolio.csv."),
                 ("Open questions", "Counted from each spec's §6, the same items handoff_questions.py writes into REPORT.md for the Agents board."),
                 ("Status", "Brownfield (live) = real code. Greenfield = README + spec. Greenfield · host = also absorbs merged features. Merged → feature = lives at <host>/features/<name>/. Parked/Retired = archived on GitHub. Built elsewhere = finished outside E:\\GIT_ROOT\\Projects."),
                 ("Fit scores", "Judgment calls (High/Medium/Low) from each README's product path, market and privacy constraints; not market research.")):
        lg.append([a, b])
    lg.column_dimensions["A"].width, lg.column_dimensions["B"].width = 22, 120
    for sheet in (mg, dg, sm, lg):
        for c in sheet[1]:
            c.fill, c.font = hdr_fill, hdr_font
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if not c.font.bold:
                    c.font = body
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(XLSX_PATH)


def write_md(rows, missing):
    def cell(v):
        return str(v or "—").replace("|", "\\|").replace("\n", " ")
    from collections import Counter
    st, use = Counter(r["status"] for r in rows), Counter(r["best_use"] for r in rows)
    active = [r for r in rows if r["status"] not in INACTIVE]
    out = ["---", "generated: by scripts/portfolio/build_portfolio.py — edit scripts/portfolio/portfolio.csv, not this file", "---", "",
           "# Project Portfolio", "",
           "Every project under `E:\\GIT_ROOT\\Projects`, with its status, best use, Dashboard Command Center group, fit as a product, "
           "and how many pre-build questions wait in its `REPORT.md` (also shown in the Agents board's Needs input panel). "
           "The same data with filters and live counts is in [[Project Portfolio.xlsx]].", "",
           f"**{len(rows)} projects** · " + " · ".join(f"{n} {k}" for k, n in sorted(st.items(), key=lambda x: -x[1])), "",
           "**Best use:** " + " · ".join(f"{n} {k}" for k, n in sorted(use.items(), key=lambda x: -x[1])), ""]
    if missing:
        out += [f"> New folders not yet in portfolio.csv: {', '.join(missing)}", ""]
    out += ["## Active projects", "",
            "| Project | Status | Best use | DCC group | Best path | Open-source | Commercial | Open questions | Purpose |",
            "|---|---|---|---|---|---|---|---|---|"]
    for r in active:
        p = r["purpose"] if len(r["purpose"]) <= 140 else r["purpose"][:139].rstrip() + "…"
        out.append("| " + " | ".join(cell(x) for x in (r["project"], r["status"], r["best_use"], r["dcc_group"], r["best_path"],
                                                        r["open_source_fit"], r["commercial_fit"], r["open"], p)) + " |")
    out += ["", "## Merged, parked, retired and excluded", "", "| Project | Status | Now lives in | Note |", "|---|---|---|---|"]
    for r in rows:
        if r["status"] in INACTIVE:
            out.append("| " + " | ".join(cell(x) for x in (r["project"], r["status"], r["merged_or_parked_into"], r["rationale"])) + " |")
    MD_PATH.write_text("\n".join(out) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
